"""
Secure Vision AI - Decision Threshold Tuning Module

Role:
    Analytics Engineer

Purpose:
    Turn model confidence scores into a business decision rule and
    choose the confidence threshold that minimises expected cost.

Decision rule:
    confidence >= threshold  -> accept the model's prediction automatically
    confidence <  threshold  -> send the image to human review

Cost model (every value is an ASSUMPTION - edit the config block below):
    - Each image sent to human review costs REVIEW_COST. Reviewers are
      assumed to give the correct label.
    - An accepted prediction that is wrong costs money. Misclassifying
      true class i as class j is treated as two errors at once:
          Type II error (missed detection of i)  -> MISSED_DETECTION_COST[i]
          Type I  error (false alarm for j)      -> FALSE_ALARM_COST[j]
      so cost[i, j] = MISSED_DETECTION_COST[i] + FALSE_ALARM_COST[j],
      and cost[i, i] = 0.

Method:
    The threshold is tuned on the validation split and then reported on
    the test split, so the reported cost is not tuned to the data it is
    measured on.

Input:
    data/processed/batch_predictions.csv (from generate_probabilities.py)
    Required columns: model, confidence, predicted_index and either
    true_index or true_class. Optional column: split.

Outputs (data/processed/threshold_tuning/):
    threshold_tuning_sheet.xlsx   the tuning sheet, with a chart
    threshold_summary.csv         optimal threshold and KPIs per model
    threshold_sweep.csv           every threshold for every model
    type_i_ii_by_class.csv        Type I / Type II breakdown per class
    cost_sensitivity.csv          optimal threshold vs review cost
    cost_matrix_dollars.csv       the dollar cost matrix used
    cost_vs_threshold.png         cost-versus-threshold chart
"""

import argparse
from pathlib import Path

import numpy as np
import pandas as pd


# ==========================================================
# PATHS
# ==========================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"
THRESHOLD_DIR = PROCESSED_DIR / "threshold_tuning"
DEFAULT_INPUT = PROCESSED_DIR / "batch_predictions.csv"


# ==========================================================
# CLASSES AND MODELS
# ==========================================================

CLASS_NAMES = [
    "airplane",
    "automobile",
    "bird",
    "cat",
    "deer",
    "dog",
    "frog",
    "horse",
    "ship",
    "truck",
]

NUM_CLASSES = len(CLASS_NAMES)

MODEL_NAMES = ("baseline_cnn", "resnet18")


# ==========================================================
# COST ASSUMPTIONS  (edit these, then document your reasons)
# ==========================================================
# These are illustrative placeholders, NOT measured financial costs.
# Example of an asymmetric setting:
#     MISSED_DETECTION_COST["truck"] = 25.0

REVIEW_COST = 1.0

MISSED_DETECTION_COST = {name: 10.0 for name in CLASS_NAMES}   # Type II
FALSE_ALARM_COST = {name: 5.0 for name in CLASS_NAMES}         # Type I


# ==========================================================
# SWEEP SETTINGS
# ==========================================================

THRESHOLDS = np.round(np.arange(0.0, 1.0, 0.01), 2)

# Review costs used in the sensitivity analysis.
SENSITIVITY_REVIEW_COSTS = (0.25, 0.5, 1.0, 2.0, 5.0, 10.0)

TUNE_SPLIT = "validation"
REPORT_SPLIT = "test"


# ==========================================================
# COST MODEL
# ==========================================================

def _cost_vector(costs, label):
    """Turn a {class_name: cost} dict into a validated array."""

    missing = set(CLASS_NAMES) - set(costs)

    if missing:
        raise ValueError(
            f"{label} is missing classes: {sorted(missing)}"
        )

    vector = np.array(
        [float(costs[name]) for name in CLASS_NAMES]
    )

    if not np.all(np.isfinite(vector)) or np.any(vector < 0):
        raise ValueError(
            f"{label} must be finite and non-negative."
        )

    return vector


def build_cost_matrix():
    """
    Build the dollar cost matrix (rows = actual, columns = predicted).

    Returns the matrix plus the per-class Type II (missed detection)
    and Type I (false alarm) cost vectors it was built from.
    """

    if not np.isfinite(REVIEW_COST) or REVIEW_COST < 0:
        raise ValueError("REVIEW_COST must be finite and non-negative.")

    missed = _cost_vector(MISSED_DETECTION_COST, "MISSED_DETECTION_COST")
    false_alarm = _cost_vector(FALSE_ALARM_COST, "FALSE_ALARM_COST")

    matrix = missed[:, None] + false_alarm[None, :]
    np.fill_diagonal(matrix, 0.0)

    return matrix, missed, false_alarm


# ==========================================================
# LOADING AND VALIDATION
# ==========================================================

def _convert_true_labels(values):
    """Accept class indices, class names, or numbers stored as text."""

    if pd.api.types.is_numeric_dtype(values):
        return values.astype(float)

    lookup = {name: i for i, name in enumerate(CLASS_NAMES)}

    cleaned = values.astype(str).str.strip().str.lower()
    indices = cleaned.map(lookup).astype(float)

    unresolved = indices.isna()

    if unresolved.any():
        indices[unresolved] = pd.to_numeric(
            cleaned[unresolved], errors="coerce"
        )

    if indices.isna().any():
        bad = values[indices.isna()].astype(str).unique().tolist()
        raise ValueError(f"Unknown true class values: {bad}")

    return indices


def load_predictions(csv_path):
    """
    Read a prediction CSV and return a clean DataFrame with columns:
    model, split, y_true, y_pred, confidence.
    """

    csv_path = Path(csv_path)

    if not csv_path.exists():
        raise FileNotFoundError(
            f"Prediction file not found: {csv_path}\n"
            "Run 'python -m src.analytics.generate_probabilities' first."
        )

    df = pd.read_csv(csv_path)

    required = {"model", "predicted_index", "confidence"}
    missing = required - set(df.columns)

    if missing:
        raise ValueError(
            "Prediction file is missing columns: "
            + ", ".join(sorted(missing))
        )

    if "true_index" in df.columns:
        true_column = "true_index"
    elif "true_class" in df.columns:
        true_column = "true_class"
    else:
        raise ValueError(
            "Prediction file needs a true_index or true_class column."
        )

    if "split" not in df.columns:
        df["split"] = "all"

    df = df.dropna(
        subset=["model", "predicted_index", "confidence", true_column]
    ).copy()

    if df.empty:
        raise ValueError("No usable rows in the prediction file.")

    clean = pd.DataFrame({
        "model": df["model"].astype(str),
        "split": df["split"].astype(str),
        "y_true": _convert_true_labels(df[true_column]),
        "y_pred": pd.to_numeric(df["predicted_index"], errors="raise"),
        "confidence": pd.to_numeric(df["confidence"], errors="raise"),
    })

    for column in ("y_true", "y_pred"):
        values = clean[column].to_numpy()

        if not np.all(np.isfinite(values)) or not np.all(
            values == np.floor(values)
        ):
            raise ValueError(f"{column} must contain whole numbers.")

        if values.min() < 0 or values.max() >= NUM_CLASSES:
            raise ValueError(
                f"{column} must be between 0 and {NUM_CLASSES - 1}."
            )

        clean[column] = values.astype(int)

    confidence = clean["confidence"].to_numpy()

    if not np.all(np.isfinite(confidence)) or np.any(
        (confidence < 0) | (confidence > 1)
    ):
        raise ValueError("confidence must be between 0 and 1.")

    return clean.reset_index(drop=True)


# ==========================================================
# THRESHOLD EVALUATION
# ==========================================================

def evaluate_threshold(
    y_true, y_pred, confidence, threshold, cost_matrix, review_cost
):
    """Apply one confidence threshold and return its metrics."""

    total = len(y_true)
    accepted = confidence >= threshold

    yt = y_true[accepted]
    yp = y_pred[accepted]

    accepted_count = int(accepted.sum())
    reviewed_count = total - accepted_count

    wrong_accepted = int(np.sum(yt != yp))
    wrong_reviewed = int(
        np.sum((~accepted) & (y_true != y_pred))
    )

    error_cost = float(cost_matrix[yt, yp].sum())
    review_total = float(reviewed_count * review_cost)
    total_cost = error_cost + review_total

    return {
        "threshold": float(threshold),
        "samples": int(total),
        "accepted": accepted_count,
        "reviewed": reviewed_count,
        "coverage": accepted_count / total,
        "review_rate": reviewed_count / total,
        "accuracy_accepted": (
            float(np.mean(yt == yp)) if accepted_count else np.nan
        ),
        "errors_accepted": wrong_accepted,
        "errors_caught_by_review": wrong_reviewed,
        "error_cost": error_cost,
        "review_cost": review_total,
        "total_cost": total_cost,
        "cost_per_1000": total_cost / total * 1000,
    }


def sweep_thresholds(
    y_true, y_pred, confidence, cost_matrix, review_cost,
    thresholds=THRESHOLDS,
):
    """Evaluate every threshold and add baseline comparisons."""

    rows = [
        evaluate_threshold(
            y_true, y_pred, confidence, t, cost_matrix, review_cost
        )
        for t in thresholds
    ]

    sweep = pd.DataFrame(rows)

    # Baseline 1: trust every prediction (no human review).
    trust_all = float(cost_matrix[y_true, y_pred].sum())

    # Baseline 2: send every image to a human.
    review_all = float(len(y_true) * review_cost)

    def savings(baseline):
        if baseline == 0:
            return np.nan
        return (baseline - sweep["total_cost"]) / baseline * 100

    sweep["savings_vs_trust_all_pct"] = savings(trust_all)
    sweep["savings_vs_review_all_pct"] = savings(review_all)

    sweep.attrs["trust_all_cost"] = trust_all
    sweep.attrs["review_all_cost"] = review_all

    return sweep


def optimal_index(sweep):
    """
    Row index of the lowest-cost threshold. Ties go to the lowest
    threshold, which keeps the most automation.
    """

    return int(sweep["total_cost"].idxmin())


def type_i_ii_by_class(
    y_true, y_pred, confidence, threshold, missed, false_alarm
):
    """
    Type I / Type II breakdown among ACCEPTED predictions.

    For class c (one-vs-rest):
        Type I  (false alarm)     = predicted c, actually not c
        Type II (missed detection) = actually c, predicted not c
    """

    accepted = confidence >= threshold
    yt = y_true[accepted]
    yp = y_pred[accepted]

    rows = []

    for index, name in enumerate(CLASS_NAMES):
        false_positives = int(np.sum((yp == index) & (yt != index)))
        false_negatives = int(np.sum((yt == index) & (yp != index)))

        rows.append({
            "class_name": name,
            "type_i_false_alarms": false_positives,
            "type_i_cost": float(false_positives * false_alarm[index]),
            "type_ii_missed_detections": false_negatives,
            "type_ii_cost": float(false_negatives * missed[index]),
        })

    result = pd.DataFrame(rows)

    result["total_error_cost"] = (
        result["type_i_cost"] + result["type_ii_cost"]
    )

    return result


def cost_sensitivity(y_true, y_pred, confidence, cost_matrix):
    """How the optimal threshold moves as the review cost changes."""

    rows = []

    for review_cost in SENSITIVITY_REVIEW_COSTS:
        sweep = sweep_thresholds(
            y_true, y_pred, confidence, cost_matrix, review_cost
        )

        best = sweep.loc[optimal_index(sweep)]

        rows.append({
            "review_cost": review_cost,
            "optimal_threshold": best["threshold"],
            "coverage": best["coverage"],
            "cost_per_1000": best["cost_per_1000"],
        })

    return pd.DataFrame(rows)


# ==========================================================
# SPLIT SELECTION
# ==========================================================

def _choose_splits(model_df, tune_split, report_split):
    """Return (tuning rows, reporting rows, note about the choice)."""

    tune = model_df[model_df["split"] == tune_split]
    report = model_df[model_df["split"] == report_split]

    if not tune.empty and not report.empty:
        return tune, report, ""

    if not report.empty:
        note = (
            f"No '{tune_split}' rows; tuned AND reported on "
            f"'{report_split}' (results are optimistic)."
        )
        return report, report, note

    if not tune.empty:
        note = (
            f"No '{report_split}' rows; tuned AND reported on "
            f"'{tune_split}' (results are optimistic)."
        )
        return tune, tune, note

    note = (
        "No validation/test split labels; tuned AND reported on all "
        "rows (results are optimistic)."
    )
    return model_df, model_df, note


def _arrays(frame):
    return (
        frame["y_true"].to_numpy(),
        frame["y_pred"].to_numpy(),
        frame["confidence"].to_numpy(),
    )


# ==========================================================
# MAIN ANALYSIS
# ==========================================================

def run_analysis(csv_path, tune_split=TUNE_SPLIT, report_split=REPORT_SPLIT):
    """Run the full threshold analysis and return all result tables."""

    df = load_predictions(csv_path)

    cost_matrix, missed, false_alarm = build_cost_matrix()

    present = set(df["model"].unique())
    models = [m for m in MODEL_NAMES if m in present] or sorted(present)

    summary_rows = []
    sweep_frames = []
    type_frames = []
    sensitivity_frames = []
    notes = []

    for model_name in models:

        model_df = df[df["model"] == model_name]

        tune_df, report_df, note = _choose_splits(
            model_df, tune_split, report_split
        )

        if note:
            notes.append(f"{model_name}: {note}")

        t_true, t_pred, t_conf = _arrays(tune_df)
        r_true, r_pred, r_conf = _arrays(report_df)

        tune_sweep = sweep_thresholds(
            t_true, t_pred, t_conf, cost_matrix, REVIEW_COST
        )

        report_sweep = sweep_thresholds(
            r_true, r_pred, r_conf, cost_matrix, REVIEW_COST
        )

        # Both sweeps use the same thresholds, so the same row index
        # is the same threshold in each.
        chosen = optimal_index(tune_sweep)
        chosen_threshold = float(tune_sweep.loc[chosen, "threshold"])

        report_row = report_sweep.loc[chosen]

        report_best = report_sweep.loc[optimal_index(report_sweep)]

        trust_all = report_sweep.attrs["trust_all_cost"]
        review_all = report_sweep.attrs["review_all_cost"]
        n = int(report_row["samples"])

        summary_rows.append({
            "model": model_name,
            "tuning_split": (
                tune_split if not tune_df.empty else "n/a"
            ),
            "reporting_split": report_split,
            "optimal_threshold": chosen_threshold,
            "samples_reported": n,
            "coverage": report_row["coverage"],
            "review_rate": report_row["review_rate"],
            "accuracy_accepted": report_row["accuracy_accepted"],
            "errors_accepted": int(report_row["errors_accepted"]),
            "errors_caught_by_review": int(
                report_row["errors_caught_by_review"]
            ),
            "total_cost": report_row["total_cost"],
            "cost_per_1000": report_row["cost_per_1000"],
            "trust_all_cost_per_1000": trust_all / n * 1000,
            "review_all_cost_per_1000": review_all / n * 1000,
            "savings_vs_trust_all_pct": report_row[
                "savings_vs_trust_all_pct"
            ],
            "savings_vs_review_all_pct": report_row[
                "savings_vs_review_all_pct"
            ],
            "best_possible_threshold_on_report_split": report_best[
                "threshold"
            ],
        })

        for role, split_name, frame in (
            ("tuning", tune_split, tune_sweep),
            ("report", report_split, report_sweep),
        ):
            labelled = frame.copy()
            labelled.insert(0, "model", model_name)
            labelled.insert(1, "role", role)
            labelled.insert(2, "split", split_name)
            sweep_frames.append(labelled)

        type_table = type_i_ii_by_class(
            r_true, r_pred, r_conf, chosen_threshold, missed, false_alarm
        )
        type_table.insert(0, "model", model_name)
        type_table.insert(1, "threshold", chosen_threshold)
        type_frames.append(type_table)

        sensitivity = cost_sensitivity(
            t_true, t_pred, t_conf, cost_matrix
        )
        sensitivity.insert(0, "model", model_name)
        sensitivity_frames.append(sensitivity)

    return {
        "summary": pd.DataFrame(summary_rows),
        "sweep": pd.concat(sweep_frames, ignore_index=True),
        "type_i_ii": pd.concat(type_frames, ignore_index=True),
        "sensitivity": pd.concat(sensitivity_frames, ignore_index=True),
        "cost_matrix": pd.DataFrame(
            cost_matrix, index=CLASS_NAMES, columns=CLASS_NAMES
        ),
        "models": models,
        "notes": notes,
    }


# ==========================================================
# DISPLAY
# ==========================================================

def print_summary(results):
    """Print a readable summary in the terminal."""

    print("\n" + "=" * 60)
    print("DECISION THRESHOLD TUNING - SUMMARY")
    print("=" * 60)

    for note in results["notes"]:
        print(f"WARNING - {note}")

    print(
        f"Review cost per image : {REVIEW_COST:.2f}\n"
        "Rule: confidence >= threshold -> automatic, "
        "otherwise human review"
    )

    for _, row in results["summary"].iterrows():
        print("\n" + "-" * 60)
        print(f"Model: {row['model']}")
        print("-" * 60)
        print(
            f"Optimal threshold (tuned on {row['tuning_split']}) : "
            f"{row['optimal_threshold']:.2f}"
        )
        print(f"Reported on                 : {row['reporting_split']}"
              f" ({row['samples_reported']} images)")
        print(f"Automated (coverage)        : {row['coverage']:.1%}")
        print(f"Sent to human review        : {row['review_rate']:.1%}")
        print(
            f"Accuracy on automated       : "
            f"{row['accuracy_accepted']:.4f}"
        )
        print(f"Wrong automated decisions   : {row['errors_accepted']}")
        print(
            f"Errors caught by review     : "
            f"{row['errors_caught_by_review']}"
        )
        print(f"Cost per 1,000 images       : {row['cost_per_1000']:.2f}")
        print(
            f"  trust everything          : "
            f"{row['trust_all_cost_per_1000']:.2f}"
        )
        print(
            f"  review everything         : "
            f"{row['review_all_cost_per_1000']:.2f}"
        )
        print(
            f"Saving vs trust everything  : "
            f"{row['savings_vs_trust_all_pct']:.1f}%"
        )
        print(
            f"Saving vs review everything : "
            f"{row['savings_vs_review_all_pct']:.1f}%"
        )


# ==========================================================
# SAVING
# ==========================================================

def _assumptions_table(results):
    rows = [
        ("Review cost per image", REVIEW_COST),
        ("Human review assumed correct", "yes"),
        ("Tuning split", TUNE_SPLIT),
        ("Reporting split", REPORT_SPLIT),
        (
            "Threshold sweep",
            f"{THRESHOLDS.min():.2f} to {THRESHOLDS.max():.2f}, "
            f"step 0.01",
        ),
        (
            "Cost rule",
            "cost[i,j] = missed_detection[i] + false_alarm[j]; "
            "cost[i,i] = 0",
        ),
        ("Status of costs", "Illustrative assumptions, not measured"),
    ]

    for name in CLASS_NAMES:
        rows.append((
            f"Missed detection (Type II) - {name}",
            MISSED_DETECTION_COST[name],
        ))

    for name in CLASS_NAMES:
        rows.append((
            f"False alarm (Type I) - {name}",
            FALSE_ALARM_COST[name],
        ))

    return pd.DataFrame(rows, columns=["assumption", "value"])


def _save_png(results, path):
    """Cost-versus-threshold chart. Skipped if matplotlib is missing."""

    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        print("matplotlib not installed - skipping PNG chart.")
        return False

    sweep = results["sweep"]
    summary = results["summary"].set_index("model")

    figure, axis = plt.subplots(figsize=(9, 5.5))

    for model_name in results["models"]:
        frame = sweep[
            (sweep["model"] == model_name) & (sweep["role"] == "report")
        ]

        line, = axis.plot(
            frame["threshold"],
            frame["cost_per_1000"],
            label=model_name,
        )

        threshold = summary.loc[model_name, "optimal_threshold"]
        cost = summary.loc[model_name, "cost_per_1000"]

        axis.scatter(
            [threshold], [cost],
            color=line.get_color(), zorder=5, s=60,
        )
        axis.annotate(
            f"t={threshold:.2f}",
            (threshold, cost),
            textcoords="offset points",
            xytext=(6, 8),
        )

    review_all = summary["review_all_cost_per_1000"].iloc[0]

    axis.axhline(
        review_all, linestyle="--", color="grey",
        label="review everything",
    )

    axis.set_xlabel("Confidence threshold")
    axis.set_ylabel("Expected cost per 1,000 images")
    axis.set_title("Expected cost versus decision threshold")
    axis.legend()
    axis.grid(alpha=0.3)

    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)

    return True


def _save_excel(results, path):
    """Tuning sheet with a native Excel chart. Skipped without openpyxl."""

    try:
        from openpyxl.chart import LineChart, Reference
    except ImportError:
        print("openpyxl not installed - skipping Excel sheet.")
        return False

    sweep = results["sweep"]

    chart_data = (
        sweep[sweep["role"] == "report"]
        .pivot(index="threshold", columns="model", values="cost_per_1000")
        [results["models"]]
        .reset_index()
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:

        results["summary"].to_excel(
            writer, sheet_name="Summary", index=False
        )

        _assumptions_table(results).to_excel(
            writer, sheet_name="Assumptions", index=False
        )

        results["cost_matrix"].to_excel(
            writer, sheet_name="Cost_Matrix"
        )

        for model_name in results["models"]:

            (
                sweep[
                    (sweep["model"] == model_name)
                    & (sweep["role"] == "report")
                ]
                .drop(columns=["model", "role"])
                .to_excel(
                    writer,
                    sheet_name=f"Sweep_{model_name}"[:31],
                    index=False,
                )
            )

            (
                results["type_i_ii"][
                    results["type_i_ii"]["model"] == model_name
                ]
                .drop(columns=["model"])
                .to_excel(
                    writer,
                    sheet_name=f"TypeI_II_{model_name}"[:31],
                    index=False,
                )
            )

        results["sensitivity"].to_excel(
            writer, sheet_name="Sensitivity", index=False
        )

        chart_data.to_excel(
            writer, sheet_name="Chart_Data", index=False
        )

        sheet = writer.sheets["Chart_Data"]
        model_count = len(results["models"])
        last_row = sheet.max_row

        chart = LineChart()
        chart.title = "Expected cost per 1,000 images vs threshold"
        chart.x_axis.title = "Confidence threshold"
        chart.y_axis.title = "Cost per 1,000 images"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.height = 10
        chart.width = 20

        chart.add_data(
            Reference(
                sheet, min_col=2, max_col=1 + model_count,
                min_row=1, max_row=last_row,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(sheet, min_col=1, min_row=2, max_row=last_row)
        )

        sheet.add_chart(chart, "F2")

        # Column widths for readability
        for name in writer.sheets:
            worksheet = writer.sheets[name]
            for column in worksheet.columns:
                width = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column
                )
                worksheet.column_dimensions[
                    column[0].column_letter
                ].width = min(max(width + 2, 10), 60)

    return True


def save_results(results):
    """Write all CSVs, the Excel sheet and the PNG chart."""

    THRESHOLD_DIR.mkdir(parents=True, exist_ok=True)

    paths = {
        "summary": THRESHOLD_DIR / "threshold_summary.csv",
        "sweep": THRESHOLD_DIR / "threshold_sweep.csv",
        "type_i_ii": THRESHOLD_DIR / "type_i_ii_by_class.csv",
        "sensitivity": THRESHOLD_DIR / "cost_sensitivity.csv",
    }

    for key, path in paths.items():
        results[key].to_csv(path, index=False)

    matrix_path = THRESHOLD_DIR / "cost_matrix_dollars.csv"
    results["cost_matrix"].to_csv(matrix_path)

    saved = list(paths.values()) + [matrix_path]

    excel_path = THRESHOLD_DIR / "threshold_tuning_sheet.xlsx"

    if _save_excel(results, excel_path):
        saved.append(excel_path)

    png_path = THRESHOLD_DIR / "cost_vs_threshold.png"

    if _save_png(results, png_path):
        saved.append(png_path)

    print("\nFiles saved:")

    for path in saved:
        print(path)

    return saved


# ==========================================================
# MAIN PROGRAM
# ==========================================================

def main():

    parser = argparse.ArgumentParser(
        description="Decision threshold tuning for Secure Vision AI."
    )
    parser.add_argument("--input", default=str(DEFAULT_INPUT))
    parser.add_argument("--tune-split", default=TUNE_SPLIT)
    parser.add_argument("--report-split", default=REPORT_SPLIT)
    args = parser.parse_args()

    print("=" * 60)
    print("Secure Vision AI - Decision Threshold Tuning")
    print("=" * 60)
    print(f"\nInput file : {args.input}")

    try:
        results = run_analysis(
            args.input, args.tune_split, args.report_split
        )
    except (ValueError, FileNotFoundError) as error:
        print(f"\nError: {error}")
        return

    print_summary(results)
    save_results(results)

    print("\n" + "=" * 60)
    print("Threshold tuning completed.")
    print("=" * 60)


if __name__ == "__main__":
    main()